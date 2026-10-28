"""Exports raster code descriptions for all six shapefiles to a Word document."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

OUT_PATH = Path("/home/arora/Documents/FSM_Slovenia/FSM_Pipeline_LULC/10_Documentation/Raster_Code_Descriptions.docx")

# ── Data ──────────────────────────────────────────────────────────────────

TABLES = [
    {
        "title": "1. Hydrogeological Units — Hgk_grouped",
        "subtitle": "Raster burn value: TYPE_CODE",
        "headers": ["TYPE_CODE", "TYPE_ENG", "Merged LAWA Sub-codes"],
        "rows": [
            ["1",  "Carbonate",                                                          "III; VI"],
            ["2",  "Carbonate/Silicate",                                                 "II"],
            ["3",  "Moderately karstified – Carbonate",                                  "X"],
            ["4",  "Moderately karstified – Organic/Carbonate",                          "XI"],
            ["5",  "Overlaying strata with essentially no groundwater resources",         "XV"],
            ["6",  "Poorly karstified – Carbonate",                                      "XII"],
            ["7",  "Silicate",                                                            "I; IV"],
            ["8",  "Silicate/Carbonate",                                                 "V"],
            ["9",  "Strata reach with organic material",                                 "XIV"],
            ["10", "Very karstified – Carbonate",                                        "VIII"],
        ],
    },
    {
        "title": "2. Lithology — lito_poly_250_d96_grouped",
        "subtitle": "Raster burn value: LITO_GRP",
        "headers": ["LITO_GRP", "Description"],
        "rows": [
            ["0",  "Quaternary loose sediments"],
            ["1",  "Organic and peat deposits"],
            ["2",  "Anthropogenic and special deposits"],
            ["3",  "Neogene clastic sediments"],
            ["4",  "Miocene carbonate rocks"],
            ["5",  "Paleogene flysch"],
            ["6",  "Cretaceous–Paleocene carbonate rocks"],
            ["7",  "Jurassic carbonate rocks"],
            ["8",  "Triassic carbonate and clastic rocks"],
            ["9",  "Paleozoic clastic and volcanic rocks"],
            ["10", "Igneous and volcanic rocks"],
            ["11", "Metamorphic rocks"],
        ],
    },
    {
        "title": "3. Soil Types — Pedoloska_grouped",
        "subtitle": "Raster burn value: SOIL_GRP",
        "headers": ["SOIL_GRP", "Description"],
        "rows": [
            ["0",  "Non-soil and artificial surfaces"],
            ["1",  "Lithosol and Regosol"],
            ["2",  "Rendzina"],
            ["3",  "Brown calcareous soils"],
            ["4",  "Ranker"],
            ["5",  "Eutric brown soils"],
            ["6",  "Dystric brown soils"],
            ["7",  "Leached and podzolic soils"],
            ["8",  "Terra rossa"],
            ["9",  "Alluvial and colluvial soils"],
            ["10", "Rigosol"],
            ["11", "Pseudogley"],
            ["12", "Gley and peat soils"],
        ],
    },
    {
        "title": "4. Rivers — Slovenia_Rivers",
        "subtitle": "Raster burn value: CODE (= SIRINA river width class)",
        "headers": ["CODE", "Width Class (SIRINA)"],
        "rows": [
            ["1",    "< 1 m"],
            ["2",    "1–2 m"],
            ["3",    "2–5 m"],
            ["4",    "5–10 m"],
            ["5",    "10–20 m"],
            ["6",    "20–50 m"],
            ["7",    "50–100 m"],
            ["8",    "> 100 m"],
            ["9999", "Unknown / unclassified"],
        ],
    },
    {
        "title": "5. Roads — Slovenia_Roads",
        "subtitle": "Raster burn value: CODE",
        "headers": ["CODE", "Road Type"],
        "rows": [
            ["0", "Local / access road"],
            ["1", "Municipal road"],
            ["2", "Regional road"],
            ["3", "Main road"],
            ["4", "Motorway / expressway"],
        ],
    },
    {
        "title": "6. Flood Extent — Flooded_NoFlood",
        "subtitle": "Raster burn value: CODE",
        "headers": ["CODE", "Class"],
        "rows": [
            ["0", "No flood"],
            ["1", "Flooded"],
        ],
    },
]

NOTE = (
    "Note: SIRINA width class breakpoints and Road Code categories follow standard "
    "Slovenian GURS/MKGP conventions — verify against source metadata if they differ. "
    "All rasters produced at 12.5 m resolution, CRS: EPSG:3794."
)

# ── Helpers ───────────────────────────────────────────────────────────────

HEADER_BG = "2E75B6"   # blue
HEADER_FG = "FFFFFF"   # white

def set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell, **kwargs):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:val"),   "single")
        tag.set(qn("w:sz"),    "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "AAAAAA")
        tcBorders.append(tag)
    tcPr.append(tcBorders)

def add_table(doc: Document, data: dict):
    # Section title
    p = doc.add_paragraph()
    run = p.add_run(data["title"])
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    # Subtitle
    sub = doc.add_paragraph()
    sr  = sub.add_run(data["subtitle"])
    sr.italic = True
    sr.font.size = Pt(9)
    sr.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
    sub.paragraph_format.space_after = Pt(4)

    headers = data["headers"]
    rows    = data["rows"]
    ncols   = len(headers)

    tbl = doc.add_table(rows=1 + len(rows), cols=ncols)
    tbl.style = "Table Grid"

    # Column widths
    col_widths = {
        1: [Cm(2.5), Cm(9),   Cm(3.5)],
        2: [Cm(2.5), Cm(11)],
    }.get(ncols, [Cm(3)] * ncols)

    for i, width in enumerate(col_widths):
        for cell in tbl.columns[i].cells:
            cell.width = width

    # Header row
    hdr_row = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size  = Pt(10)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_bg(cell, HEADER_BG)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = tbl.rows[r_idx + 1]
        bg  = "F2F7FC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = val
            para = cell.paragraphs[0]
            para.runs[0].font.size = Pt(10)
            if c_idx == 0:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.runs[0].bold = True
            set_cell_bg(cell, bg)

    doc.add_paragraph()   # spacing after table


# ── Build document ────────────────────────────────────────────────────────

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# Document title
title_p = doc.add_heading("Raster Code Descriptions", level=0)
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

sub_p = doc.add_paragraph("FSM Slovenia — Vector-to-Raster Attribute Codebook")
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.runs[0].italic = True
sub_p.runs[0].font.size = Pt(11)
sub_p.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)
doc.add_paragraph()

for data in TABLES:
    add_table(doc, data)

# Footer note
note_p = doc.add_paragraph()
note_r = note_p.add_run(NOTE)
note_r.italic = True
note_r.font.size = Pt(9)
note_r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

doc.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
